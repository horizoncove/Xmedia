#!/usr/bin/env python3
"""Build XLOOP PARK AI 5W principles PowerPoint."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Visual system: night-city commercial — deep ink + warm amber + soft sand
INK = RGBColor(0x14, 0x16, 0x1A)
SLATE = RGBColor(0x2C, 0x33, 0x3A)
MUTED = RGBColor(0x8A, 0x91, 0x99)
SAND = RGBColor(0xF3, 0xEE, 0xE6)
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
AMBER = RGBColor(0xD4, 0x8A, 0x3A)
AMBER_D = RGBColor(0xB8, 0x6E, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xDB, 0xD0)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
TOTAL = 14


def set_run(run, size=18, bold=False, color=INK, font="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", "Microsoft YaHei")


def add_rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=18,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        anchor_map = {
            MSO_ANCHOR.TOP: "t",
            MSO_ANCHOR.MIDDLE: "ctr",
            MSO_ANCHOR.BOTTOM: "b",
        }
        tf._txBody.bodyPr.set("anchor", anchor_map[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def new_slide(bg=CREAM):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, W, H, bg)
    return slide


def footer(slide, page):
    add_text(
        slide,
        Inches(0.6),
        Inches(7.1),
        Inches(8),
        Inches(0.3),
        "XLOOP PARK  ·  AI 5W 原则  ·  战略认知简报",
        size=10,
        color=MUTED,
    )
    add_text(
        slide,
        Inches(11.5),
        Inches(7.1),
        Inches(1.3),
        Inches(0.3),
        f"{page} / {TOTAL}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def section_label(slide, text, y=Inches(0.45)):
    add_rect(slide, Inches(0.6), y + Inches(0.08), Inches(0.18), Inches(0.18), AMBER)
    add_text(slide, Inches(0.95), y, Inches(10), Inches(0.35), text, size=12, bold=True, color=AMBER_D)


# 1 Cover
s = new_slide(INK)
add_rect(s, 0, 0, Inches(0.18), H, AMBER)
add_text(
    s,
    Inches(0.9),
    Inches(1.8),
    Inches(11),
    Inches(0.4),
    "XLOOP PARK  ·  AI STRATEGIC BRIEF",
    size=14,
    bold=True,
    color=AMBER,
)
add_text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.2), "AI 5W 原则", size=54, bold=True, color=WHITE)
add_text(
    s,
    Inches(0.9),
    Inches(3.6),
    Inches(11.5),
    Inches(0.7),
    "如何改写商业街区的经营底层逻辑",
    size=26,
    color=SAND,
)
add_text(
    s,
    Inches(0.9),
    Inches(4.6),
    Inches(11),
    Inches(0.5),
    "Why · Who · What · When · Where",
    size=18,
    color=AMBER,
)
add_text(
    s,
    Inches(0.9),
    Inches(6.3),
    Inches(10),
    Inches(0.4),
    "基于白皮书 v0.6  ·  不是商业街区做 AI，而是 AI 原生街区做商业",
    size=13,
    color=MUTED,
)

# 2 Thesis
s = new_slide()
footer(s, 2)
section_label(s, "CORE THESIS")
add_text(
    s,
    Inches(0.6),
    Inches(0.95),
    Inches(12),
    Inches(0.7),
    "真正难的不是选供应商，而是改写五个基本问题",
    size=28,
    bold=True,
    color=INK,
)
add_text(
    s,
    Inches(0.6),
    Inches(1.8),
    Inches(12),
    Inches(0.6),
    "AI 不是给街区加工具，而是改写街区回答 5W 的方式。\n回答方式一变，商业模式、组织、投资优先级全部跟着变。",
    size=18,
    color=SLATE,
)
cards = [
    ("不是", "堆场景清单\n买大屏大脑\n对标抄作业"),
    ("而是", "重写 Why 使命\n重定义 Who 主体\n产品化 What 决策"),
    ("结果", "不动产项目\n→\n可学习的操作系统"),
]
for i, (title, body) in enumerate(cards):
    x = Inches(0.6) + i * Inches(4.1)
    add_rect(s, x, Inches(3.0), Inches(3.85), Inches(3.4), WHITE, LINE)
    add_rect(s, x, Inches(3.0), Inches(3.85), Inches(0.12), AMBER if i else SLATE)
    add_text(s, x + Inches(0.3), Inches(3.35), Inches(3.2), Inches(0.4), title, size=14, bold=True, color=AMBER)
    add_text(s, x + Inches(0.3), Inches(3.9), Inches(3.2), Inches(2.2), body, size=20, bold=True, color=INK)

# 3 Overview
s = new_slide()
footer(s, 3)
section_label(s, "5W OVERVIEW")
add_text(
    s,
    Inches(0.6),
    Inches(0.95),
    Inches(12),
    Inches(0.5),
    "五个问题：传统默认答案 vs AI 改写后",
    size=26,
    bold=True,
    color=INK,
)
rows = [
    ("Why", "收租金、聚人气", "降低错配成本，提高注意力→交易转化"),
    ("Who", "房东、大牌、路过客流", "样本型客群 + 商户节点 + 编排者 + 内容产能"),
    ("What", "铺位、空间、活动", "匹配 · 感知 · 内容 · 预测 · 归因"),
    ("When", "节假日与经验排期", "实时 · 节律 · 预测 · 开业窗口"),
    ("Where", "物理铺位与中庭", "物理场 + 内容场 + 数据场叠合"),
]
add_rect(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(0.45), INK)
add_text(s, Inches(0.8), Inches(1.78), Inches(1.5), Inches(0.35), "5W", size=12, bold=True, color=AMBER)
add_text(s, Inches(2.5), Inches(1.78), Inches(4.5), Inches(0.35), "传统街区默认答案", size=12, bold=True, color=WHITE)
add_text(s, Inches(7.2), Inches(1.78), Inches(5.2), Inches(0.35), "AI 介入后被改写为", size=12, bold=True, color=WHITE)
for i, (w_label, a, b) in enumerate(rows):
    y = Inches(2.2) + i * Inches(0.85)
    bg = WHITE if i % 2 == 0 else SAND
    add_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.8), bg, LINE)
    add_text(s, Inches(0.8), y + Inches(0.22), Inches(1.5), Inches(0.4), w_label, size=18, bold=True, color=AMBER_D)
    add_text(s, Inches(2.5), y + Inches(0.22), Inches(4.5), Inches(0.45), a, size=14, color=MUTED)
    add_text(s, Inches(7.2), y + Inches(0.22), Inches(5.2), Inches(0.45), b, size=14, bold=True, color=INK)

# 4 Causal chain
s = new_slide()
footer(s, 4)
section_label(s, "CAUSAL CHAIN")
add_text(
    s,
    Inches(0.6),
    Inches(0.95),
    Inches(12),
    Inches(0.5),
    "5W 不是并列清单，而是一条因果链",
    size=26,
    bold=True,
    color=INK,
)
chain = [
    ("Why", "为何降错配"),
    ("Who", "对谁 / 由谁"),
    ("What", "用什么能力"),
    ("When", "何时最有杠杆"),
    ("Where", "落在哪一界面"),
]
for i, (w_label, sub) in enumerate(chain):
    x = Inches(0.5) + i * Inches(2.5)
    if i == 0:
        add_rect(s, x, Inches(2.0), Inches(2.2), Inches(1.8), INK)
        add_text(s, x, Inches(2.35), Inches(2.2), Inches(0.5), w_label, size=22, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(2.95), Inches(2.0), Inches(0.6), sub, size=13, color=SAND, align=PP_ALIGN.CENTER)
    else:
        add_rect(s, x, Inches(2.0), Inches(2.2), Inches(1.8), WHITE, LINE)
        add_text(s, x, Inches(2.35), Inches(2.2), Inches(0.5), w_label, size=22, bold=True, color=INK, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.1), Inches(2.95), Inches(2.0), Inches(0.6), sub, size=13, color=SLATE, align=PP_ALIGN.CENTER)
    if i < 4:
        add_text(
            s,
            x + Inches(2.05),
            Inches(2.55),
            Inches(0.4),
            Inches(0.4),
            "→",
            size=20,
            bold=True,
            color=AMBER,
            align=PP_ALIGN.CENTER,
        )
notes = [
    ("缺 Why", "What 沦为功能堆砌"),
    ("缺 Who", "系统没有执行主体"),
    ("缺 When / Where", "模型停在报表，调不动现场"),
]
for i, (title, body) in enumerate(notes):
    x = Inches(0.6) + i * Inches(4.1)
    add_rect(s, x, Inches(4.4), Inches(3.9), Inches(1.9), SAND)
    add_text(s, x + Inches(0.3), Inches(4.65), Inches(3.3), Inches(0.4), title, size=16, bold=True, color=AMBER_D)
    add_text(s, x + Inches(0.3), Inches(5.2), Inches(3.3), Inches(0.8), body, size=16, color=INK)

# 5 Why
s = new_slide()
footer(s, 5)
section_label(s, "01  ·  WHY")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "从贩卖稀缺位置，到经营稀缺匹配", size=26, bold=True, color=INK)
add_rect(s, Inches(0.6), Inches(1.7), Inches(5.9), Inches(2.2), WHITE, LINE)
add_text(s, Inches(0.9), Inches(1.9), Inches(5.3), Inches(0.35), "传统 Why", size=14, bold=True, color=MUTED)
add_text(
    s,
    Inches(0.9),
    Inches(2.4),
    Inches(5.3),
    Inches(1.2),
    "控制好地段 → 收人流租金\n热闹 = 成功\n铺满 = 组合正确",
    size=16,
    color=INK,
)
add_rect(s, Inches(6.8), Inches(1.7), Inches(5.9), Inches(2.2), INK)
add_text(s, Inches(7.1), Inches(1.9), Inches(5.3), Inches(0.35), "AI Why", size=14, bold=True, color=AMBER)
add_text(
    s,
    Inches(7.1),
    Inches(2.4),
    Inches(5.3),
    Inches(1.2),
    "持续降低系统性错配\n提高注意力 → 交易 → 复购概率\n错配成本成为可经营变量",
    size=16,
    color=WHITE,
)
mismatches = [
    ("人—场", "来了逛不对"),
    ("人—货", "看见非想要"),
    ("货—场", "品牌放错位"),
    ("时—需", "潮汐失衡"),
    ("内—外", "种草到不了店"),
]
for i, (a, b) in enumerate(mismatches):
    x = Inches(0.6) + i * Inches(2.45)
    add_rect(s, x, Inches(4.3), Inches(2.3), Inches(2.0), SAND)
    add_text(s, x + Inches(0.15), Inches(4.55), Inches(2.0), Inches(0.4), a, size=16, bold=True, color=AMBER_D, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.15), Inches(5.15), Inches(2.0), Inches(0.7), b, size=14, color=INK, align=PP_ALIGN.CENTER)

# 6 Who
s = new_slide()
footer(s, 6)
section_label(s, "02  ·  WHO")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "谁算数：主体被重新定义", size=26, bold=True, color=INK)
whos = [
    ("客群", "匿名流量", "可学习样本", "高频师生 = 研发样本库\n不是宣传口径里的人流"),
    ("商户", "交租乙方", "网络节点", "必须同时是\n数据 / 内容 / 体验节点"),
    ("运营", "收租与物业", "编排者", "掌握匹配、触达、\n调度、归因规则"),
    ("MCN", "营销外包", "生产资料", "与铺位平行的\n内容产能所有者"),
]
for i, (role, old, new, note) in enumerate(whos):
    x = Inches(0.45) + i * Inches(3.2)
    add_rect(s, x, Inches(1.7), Inches(3.0), Inches(4.7), WHITE, LINE)
    add_rect(s, x, Inches(1.7), Inches(3.0), Inches(0.55), INK)
    add_text(s, x, Inches(1.8), Inches(3.0), Inches(0.4), role, size=16, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.5), Inches(2.6), Inches(0.3), "过去", size=11, color=MUTED)
    add_text(s, x + Inches(0.2), Inches(2.85), Inches(2.6), Inches(0.5), old, size=15, color=MUTED)
    add_text(s, x + Inches(0.2), Inches(3.5), Inches(2.6), Inches(0.3), "现在", size=11, bold=True, color=AMBER_D)
    add_text(s, x + Inches(0.2), Inches(3.85), Inches(2.6), Inches(0.5), new, size=18, bold=True, color=INK)
    add_text(s, x + Inches(0.2), Inches(4.6), Inches(2.6), Inches(1.4), note, size=13, color=SLATE)

# 7 What
s = new_slide()
footer(s, 7)
section_label(s, "03  ·  WHAT")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "产出什么：决策本身成为产品", size=26, bold=True, color=INK)
add_text(s, Inches(0.6), Inches(1.55), Inches(12), Inches(0.4), "租金仍在，但增量来自可产品化的经营能力", size=16, color=SLATE)
whats = [
    ("匹配", "对的人 × 对的供给", "转化 / 复购"),
    ("感知", "客流 · 能耗 · 情绪可见", "调度 / 降本"),
    ("内容", "可传播的故事与货盘", "到店 / 联单"),
    ("预测", "需求 · 峰值 · 存活提前量", "减少救火"),
    ("归因", "说清哪次动作真赚了", "预算可迭代"),
]
for i, (title, a, b) in enumerate(whats):
    y = Inches(2.15) + i * Inches(0.85)
    add_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.75), WHITE if i % 2 == 0 else SAND, LINE)
    add_rect(s, Inches(0.6), y, Inches(0.12), Inches(0.75), AMBER)
    add_text(s, Inches(1.0), y + Inches(0.18), Inches(2.2), Inches(0.4), title, size=18, bold=True, color=INK)
    add_text(s, Inches(3.5), y + Inches(0.18), Inches(5.5), Inches(0.4), a, size=15, color=SLATE)
    add_text(s, Inches(9.2), y + Inches(0.18), Inches(3.2), Inches(0.4), b, size=15, bold=True, color=AMBER_D)

# 8 When
s = new_slide()
footer(s, 8)
section_label(s, "04  ·  WHEN")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "时间是权力：从日历治理到可干预变量", size=26, bold=True, color=INK)
times = [
    ("实时", "秒—分钟级", "排队 · 热力 · 安保\n空调 · LED"),
    ("节律", "日 / 周 / 学期", "反潮汐：师生空窗\n用酒店与内容补"),
    ("预测", "提前准备", "活动备货 · 摊位\n定价 · 排班"),
    ("窗口", "不可逆时点", "开业前预埋\n首年心智 · 合同"),
]
for i, (title, sub, body) in enumerate(times):
    x = Inches(0.5) + i * Inches(3.2)
    if i == 3:
        add_rect(s, x, Inches(1.8), Inches(3.0), Inches(3.6), INK)
        add_text(s, x + Inches(0.25), Inches(2.1), Inches(2.5), Inches(0.4), title, size=22, bold=True, color=AMBER)
        add_text(s, x + Inches(0.25), Inches(2.65), Inches(2.5), Inches(0.35), sub, size=13, color=WHITE)
        add_text(s, x + Inches(0.25), Inches(3.3), Inches(2.5), Inches(1.6), body, size=15, color=SAND)
    else:
        add_rect(s, x, Inches(1.8), Inches(3.0), Inches(3.6), WHITE, LINE)
        add_text(s, x + Inches(0.25), Inches(2.1), Inches(2.5), Inches(0.4), title, size=22, bold=True, color=AMBER_D)
        add_text(s, x + Inches(0.25), Inches(2.65), Inches(2.5), Inches(0.35), sub, size=13, color=MUTED)
        add_text(s, x + Inches(0.25), Inches(3.3), Inches(2.5), Inches(1.6), body, size=15, color=SLATE)
add_rect(s, Inches(0.6), Inches(5.7), Inches(12.1), Inches(1.0), SAND)
add_text(
    s,
    Inches(0.9),
    Inches(5.95),
    Inches(11.5),
    Inches(0.55),
    "最高杠杆的 When：开业前约 9 个月。错过这个窗口，以后每一分钱都在付「改造税」与「组织重训税」。",
    size=15,
    bold=True,
    color=INK,
)

# 9 Where
s = new_slide()
footer(s, 9)
section_label(s, "05  ·  WHERE")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "空间从容器，变成可调度的三场", size=26, bold=True, color=INK)
fields = [
    ("物理场", "动线 · 铺位 · 广场\n停车 · 夜市", "可走、可停、可交易的实体界面"),
    ("内容场", "直播间 · LED\n短视频场景 · 数字人", "决策前移发生的注意力界面"),
    ("数据场", "热力 · 画像 · 交易\n能耗 · 工单", "让另两场可计算、可回流"),
]
for i, (title, a, b) in enumerate(fields):
    x = Inches(0.5) + i * Inches(4.2)
    add_rect(s, x, Inches(1.75), Inches(4.0), Inches(3.5), WHITE, LINE)
    add_rect(s, x, Inches(1.75), Inches(4.0), Inches(0.7), INK)
    add_text(s, x, Inches(1.9), Inches(4.0), Inches(0.45), title, size=20, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.3), Inches(2.8), Inches(3.4), Inches(1.2), a, size=16, color=INK)
    add_text(s, x + Inches(0.3), Inches(4.2), Inches(3.4), Inches(0.8), b, size=13, color=SLATE)
add_text(
    s,
    Inches(0.6),
    Inches(5.6),
    Inches(12),
    Inches(0.9),
    "分裂即失败：物理场热闹但内容场在别处带货；数据场有热力但调不了铺；内容场爆了但核销接不住。\n三场必须被同一套编排逻辑看见——夜经济主场（广场+夜市+LED）应与一楼同等系统优先级。",
    size=14,
    color=SLATE,
)

# 10 Structural shifts
s = new_slide()
footer(s, 10)
section_label(s, "STRUCTURAL SHIFT")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "5W 同向改写后的六种结构性质变", size=26, bold=True, color=INK)
shifts = [
    ("商业模式", "租售空间 → 经营匹配"),
    ("竞争壁垒", "地段 → 数据飞轮 × 内容产能 × 预埋时点"),
    ("组织权力", "经验权威 → 编排权威"),
    ("消费旅程", "到店才开始 → 来之前已被计算"),
    ("空间经济", "固定租金梯度 → 动态热力价值"),
    ("风险形态", "租不租得出 → 学不学得会"),
]
for i, (a, b) in enumerate(shifts):
    col, row = i % 3, i // 3
    x = Inches(0.5) + col * Inches(4.2)
    y = Inches(1.8) + row * Inches(2.3)
    add_rect(s, x, y, Inches(4.0), Inches(2.05), WHITE, LINE)
    add_text(s, x + Inches(0.3), y + Inches(0.35), Inches(3.4), Inches(0.4), a, size=14, bold=True, color=AMBER_D)
    add_text(s, x + Inches(0.3), y + Inches(0.9), Inches(3.4), Inches(0.9), b, size=18, bold=True, color=INK)

# 11 XLOOP endowments
s = new_slide()
footer(s, 11)
section_label(s, "XLOOP ENDOWMENT")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "不是学谁：认清自己在五个 W 上的禀赋", size=26, bold=True, color=INK)
rows = [
    ("Why", "产住商娱教一体，错配类型齐全", "使命=降错配的原生街区"),
    ("Who", "师生样本 + 居民 + 酒店 + 驻场 MCN", "主体操作系统先天完整"),
    ("What", "Xmedia 基因：内容/AIGC 非外购", "内容可做成中台级产品"),
    ("When", "距开业约 9 个月预埋窗口", "最高杠杆在现在"),
    ("Where", "MCN 楼层 + 夜经济主场已成型 7–8 成", "三场叠合物理前提部分具备"),
]
add_rect(s, Inches(0.6), Inches(1.65), Inches(12.1), Inches(0.45), INK)
add_text(s, Inches(0.85), Inches(1.73), Inches(1.5), Inches(0.3), "5W", size=12, bold=True, color=AMBER)
add_text(s, Inches(2.5), Inches(1.73), Inches(5.5), Inches(0.3), "禀赋", size=12, bold=True, color=WHITE)
add_text(s, Inches(8.2), Inches(1.73), Inches(4.2), Inches(0.3), "战略含义", size=12, bold=True, color=WHITE)
for i, (w_label, a, b) in enumerate(rows):
    y = Inches(2.15) + i * Inches(0.85)
    add_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.8), WHITE if i % 2 == 0 else SAND, LINE)
    add_text(s, Inches(0.85), y + Inches(0.22), Inches(1.5), Inches(0.4), w_label, size=16, bold=True, color=AMBER_D)
    add_text(s, Inches(2.5), y + Inches(0.22), Inches(5.5), Inches(0.45), a, size=14, color=INK)
    add_text(s, Inches(8.2), y + Inches(0.22), Inches(4.2), Inches(0.45), b, size=14, bold=True, color=SLATE)

# 12 Priority
s = new_slide()
footer(s, 12)
section_label(s, "PRIORITY BY 5W")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "用 5W 因果重排优先级（原理序）", size=24, bold=True, color=INK)
prios = [
    ("P0", "Why+Who 基建", "身份 / 会员 / POS / 商户数据义务"),
    ("P0", "What 现金牛", "内容核销 · 能耗 · 客流感知"),
    ("P0", "When 窗口", "招商 AI · MCN 预招商 · 点位预埋"),
    ("P0", "Where 主场", "夜经济指挥链路（广场/摊位/LED/安保）"),
    ("P1", "What 体验", "停车 / 点餐 / 导览（抬客单与停留）"),
    ("P2", "What 大脑", "沙盘与复杂决策（需长数据后置）"),
]
for i, (p, title, body) in enumerate(prios):
    y = Inches(1.65) + i * Inches(0.8)
    if p == "P0":
        add_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.7), INK)
        add_text(s, Inches(0.85), y + Inches(0.18), Inches(1.0), Inches(0.4), p, size=16, bold=True, color=AMBER)
        add_text(s, Inches(2.1), y + Inches(0.18), Inches(3.5), Inches(0.4), title, size=16, bold=True, color=WHITE)
        add_text(s, Inches(5.8), y + Inches(0.18), Inches(6.6), Inches(0.4), body, size=14, color=SAND)
    else:
        add_rect(s, Inches(0.6), y, Inches(12.1), Inches(0.7), WHITE, LINE)
        add_text(s, Inches(0.85), y + Inches(0.18), Inches(1.0), Inches(0.4), p, size=16, bold=True, color=AMBER_D)
        add_text(s, Inches(2.1), y + Inches(0.18), Inches(3.5), Inches(0.4), title, size=16, bold=True, color=INK)
        add_text(s, Inches(5.8), y + Inches(0.18), Inches(6.6), Inches(0.4), body, size=14, color=SLATE)

# 13 Board questions
s = new_slide()
footer(s, 13)
section_label(s, "BOARD QUESTIONS")
add_text(s, Inches(0.6), Inches(0.95), Inches(12), Inches(0.5), "董事会三问（5W 版）", size=26, bold=True, color=INK)
qs = [
    ("01", "Why", "是否同意：AI 的第一性原理是降错配，而不是上功能？"),
    ("02", "Who", "是否同意：商户与 MCN 必须被定义为节点/产能，而非纯承租方？"),
    ("03", "When", "是否同意：九个月窗口上的预埋，优先级高于开业后的「好看大脑」？"),
]
for i, (n, w_label, q) in enumerate(qs):
    y = Inches(1.8) + i * Inches(1.5)
    add_rect(s, Inches(0.6), y, Inches(12.1), Inches(1.3), WHITE, LINE)
    add_rect(s, Inches(0.6), y, Inches(1.4), Inches(1.3), INK)
    add_text(s, Inches(0.6), y + Inches(0.25), Inches(1.4), Inches(0.35), n, size=20, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.6), y + Inches(0.7), Inches(1.4), Inches(0.35), w_label, size=14, color=SAND, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.3), y + Inches(0.4), Inches(10), Inches(0.6), q, size=18, bold=True, color=INK)

# 14 Close
s = new_slide(INK)
add_rect(s, 0, 0, Inches(0.18), H, AMBER)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.4), "CLOSING", size=14, bold=True, color=AMBER)
add_text(s, Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.0), "五个 W 同向改写", size=36, bold=True, color=WHITE)
add_text(
    s,
    Inches(0.9),
    Inches(3.2),
    Inches(11.5),
    Inches(1.2),
    "街区就从「不动产项目」\n跃迁为「可学习的商业操作系统」。",
    size=22,
    color=SAND,
)
add_text(
    s,
    Inches(0.9),
    Inches(4.8),
    Inches(11.5),
    Inches(0.8),
    "XLOOP 现在要做的，不是寻找可模仿的对象，\n而是在开业前把这五个 W 一次性写进基因——把基因变成肌肉。",
    size=16,
    color=MUTED,
)
add_text(s, Inches(0.9), Inches(6.3), Inches(11), Inches(0.4), "Live · Loop · Life", size=14, bold=True, color=AMBER)

out = "/workspace/docs/ppt/XLOOP_PARK_AI_5W原则战略简报.pptx"
prs.save(out)
print("saved", out)
print("slides", len(prs.slides))
